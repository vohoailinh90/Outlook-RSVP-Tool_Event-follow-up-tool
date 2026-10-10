"""The Excel reports the app writes and the gift list it reads back.

Moved out of RSVPApp so they can be tested: these files are exported and
attached to emails sent to colleagues (the attendance report rides on the
Thank-you email, the contributor list on the gift report), and a wrong
figure in them is a wrong figure in someone's inbox.

Rosters are the dicts the app keeps per email address:
    attendance: {"name", "vote", "actual_attend", "free", "amount",
                 "extra_attends", "extra_amounts"}  (per payment round)
    gift:       {"name", "checked", "amount", "send_email"}
"""
from __future__ import annotations

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from rsvp.domain import parse_amount_from_text, round_totals


def write_header_row(ws, headers):
    """Row 1 in the app's report style: white bold text on dark blue."""
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="003366")
        c.alignment = Alignment(horizontal="center")


def attendance_workbook(roster, figures):
    """The Attendance & Payment report (Tab 5): one row per Yes/Maybe voter
    with an Attend and an amount column for every payment round, then per
    round its attendees, collected, paid and remaining, and the collected
    and remaining totals over all rounds. Used by "📊 Export to Excel" and
    as the Thank-you email's attachment. Layout as in the other lineage.

    figures: rsvp.domain.payment_rounds(...) for this roster - round 1
    first. Remaining comes from those figures, never from a label on
    screen, so an overpaid round keeps its minus sign."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Attendance & Payment"
    first, extra = figures[0], figures[1:]
    headers = ["No.", "Name", "Email", "Vote", f"{first.label} Attend", "Free", first.label]
    for f in extra:
        headers += [f"{f.label} Attend", f.label]
    write_header_row(ws, headers)
    row_idx = 2
    for i, (email, info) in enumerate(roster.items(), start=1):
        ws.cell(row=row_idx, column=1, value=i)
        ws.cell(row=row_idx, column=2, value=info["name"])
        ws.cell(row=row_idx, column=3, value=email)
        ws.cell(row=row_idx, column=4, value=info.get("vote", ""))
        ws.cell(row=row_idx, column=5, value=(info.get("actual_attend") or "").strip())
        ws.cell(row=row_idx, column=6, value="Yes" if info.get("free") else "No")
        ws.cell(row=row_idx, column=7, value=info.get("amount", 0.0) or None)
        col = 8
        for f in extra:
            attend = ((info.get("extra_attends") or {}).get(f.key) or "").strip()
            ws.cell(row=row_idx, column=col, value=attend or None)
            ws.cell(row=row_idx, column=col + 1,
                    value=(info.get("extra_amounts") or {}).get(f.key, 0.0) or None)
            col += 2
        row_idx += 1

    def total(row, label, column, value, money=True):
        ws.cell(row=row, column=2, value=label).font = Font(bold=True)
        cell = ws.cell(row=row, column=column, value=value)
        cell.font = Font(bold=True)
        if money:
            cell.number_format = "#,##0"

    r = row_idx + 1
    total(r, "Total actual attend:", 5, first.attendees, money=False)
    r += 1
    for f in figures:
        total(r, f"{f.label} — Attendees:", 5, f.attendees, money=False)
        total(r + 1, f"{f.label} — Collect amount:", 7, f.collected)
        total(r + 2, f"{f.label} — Amount paid:", 7, f.paid)
        total(r + 3, f"{f.label} — Remaining amount:", 7, f.remaining)
        r += 4
    collected, _paid, remaining = round_totals(figures)
    total(r, "Total collected amount (all rounds):", 7, collected)
    total(r + 1, "Total remaining amount (all rounds):", 7, remaining)

    for letter, width in zip("ABCDEFG", (6, 32, 30, 10, 14, 10, 14)):
        ws.column_dimensions[letter].width = width
    for i in range(8, 8 + 2 * len(extra)):
        ws.column_dimensions[get_column_letter(i)].width = 14
    return wb


def gift_contribution_workbook(roster):
    """The full gift list (Tab 6 "📊 Export to Excel"): everyone, with a
    Contributed Yes/No column, and a TOTAL COLLECTED row of the contributors'
    amounts. read_gift_contribution_rows() reads it back."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Gift Contribution"
    write_header_row(ws, ["No.", "Name", "Email", "Amount", "Contributed"])
    row_idx = 2
    total_amount = 0.0
    for i, (email, info) in enumerate(roster.items(), start=1):
        amount = info.get("amount", 0.0)
        if info["checked"]:
            total_amount += amount
        ws.cell(row=row_idx, column=1, value=i)
        ws.cell(row=row_idx, column=2, value=info["name"])
        ws.cell(row=row_idx, column=3, value=email)
        ws.cell(row=row_idx, column=4, value=amount or None)
        ws.cell(row=row_idx, column=5, value="Yes" if info["checked"] else "No")
        row_idx += 1

    # Grand total row, one blank row below the table. The label goes in the
    # NAME column (not Email) so read_gift_contribution_rows() skips this row
    # on reload: it keys off the Email column being non-empty.
    total_row = row_idx + 1
    label_cell = ws.cell(row=total_row, column=2, value="TOTAL COLLECTED:")
    label_cell.font = Font(bold=True)
    label_cell.alignment = Alignment(horizontal="right")
    total_cell = ws.cell(row=total_row, column=4, value=total_amount)
    total_cell.font = Font(bold=True)
    total_cell.number_format = "#,##0"

    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 32
    ws.column_dimensions["D"].width = 14
    ws.column_dimensions["E"].width = 14
    return wb


def gift_report_workbook(roster):
    """The attachment of the gift report email: only the people who HAVE
    contributed, renumbered from 1, as No./Name/Email/Amount - the "Send
    email"/"Contributed" checkboxes are UI state, not report data."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Contributors"
    write_header_row(ws, ["No.", "Name", "Email", "Amount"])
    row_idx = 2
    seq = 0
    total_amount = 0.0
    for email, info in roster.items():
        if not info.get("checked"):
            continue
        seq += 1
        amount = info.get("amount", 0.0)
        total_amount += amount
        ws.cell(row=row_idx, column=1, value=seq)
        ws.cell(row=row_idx, column=2, value=info.get("name") or email)
        ws.cell(row=row_idx, column=3, value=email)
        ws.cell(row=row_idx, column=4, value=amount or None)
        row_idx += 1
    total_row = row_idx + 1
    ws.cell(row=total_row, column=2, value="TOTAL COLLECTED:").font = Font(bold=True)
    c_amt = ws.cell(row=total_row, column=4, value=total_amount)
    c_amt.font = Font(bold=True)
    c_amt.number_format = "#,##0"
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 32
    ws.column_dimensions["D"].width = 14
    return wb


def read_gift_contribution_rows(path):
    """Reads a Gift_Contribution_List_*.xlsx file and returns
    list[dict(name, email, checked, amount)], matching columns by their
    HEADER text (row 1) instead of a fixed position — so both older files
    (Name/Email/Contributed only) and newer files (No./Name/Email/Amount/
    Contributed) can be read back correctly."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    header_row = next(ws.iter_rows(min_row=1, max_row=1), ())
    header = {(str(c.value).strip() if c.value else ""): i for i, c in enumerate(header_row)}
    i_name = header.get("Name")
    i_email = header.get("Email")
    i_amount = header.get("Amount")
    i_contrib = header.get("Contributed")
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row:
            continue
        email = row[i_email] if i_email is not None and i_email < len(row) else None
        if not email:
            continue  # skips blank rows and the "TOTAL COLLECTED" summary row
        name = row[i_name] if i_name is not None and i_name < len(row) else None
        contributed = row[i_contrib] if i_contrib is not None and i_contrib < len(row) else None
        amount = row[i_amount] if i_amount is not None and i_amount < len(row) else None
        rows.append({
            "name": str(name).strip() if name else str(email).strip(),
            "email": str(email).strip(),
            "checked": (str(contributed).strip().lower() == "yes") if contributed is not None else False,
            "amount": parse_amount_from_text(amount) if amount is not None else 0.0,
        })
    return rows
